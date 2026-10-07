# -*- coding: utf-8 -*-
"""md2docx — Markdown to Word (.docx) converter.

Preserves the source text as-is, only the format is converted:
headings, tables (including <br> line breaks inside cells), fenced code
blocks, blockquotes, bullet/numbered lists (numbering restarts per list),
bold and inline code, horizontal rules.

Usage:
    md2docx.py file.md                       -> file.docx (next to source)
    md2docx.py a.md b.md c.md                -> a.docx b.docx c.docx
    md2docx.py docs\\*.md                     -> every .md in docs\\
    md2docx.py --all                         -> every .md in current folder
    md2docx.py --all -o out                  -> results go to folder 'out'
    md2docx.py file.md -o result.docx        -> explicit output name

Run with the bundled venv Python (or any Python that has python-docx):
    .\\.venv\\Scripts\\python.exe md2docx.py file.md
"""
import argparse
import glob
import os
import re
import sys

try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
except ImportError:  # pragma: no cover
    sys.stderr.write(
        "python-docx is not installed for this interpreter.\n"
        "Either run:  .\\.venv\\Scripts\\python.exe md2docx.py ...\n"
        "or install:  python -m pip install python-docx\n"
    )
    sys.exit(2)

__version__ = "0.1.0"

CODE_FONT = "Consolas"
CODE_SIZE = Pt(9)
TABLE_HEADER_FILL = "D9E2F3"
CODE_FILL = "F2F2F2"


# ---------------------------------------------------------------------------
# Low-level OOXML helpers
# ---------------------------------------------------------------------------
def _shade(element, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def shade_paragraph(paragraph, fill):
    _shade(paragraph._p.get_or_add_pPr(), fill)


def shade_cell(cell, fill):
    _shade(cell._tc.get_or_add_tcPr(), fill)


def add_bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def create_numbering(doc, abstract_id=7, start=1):
    """Register a fresh numId so every ordered list restarts at `start`."""
    numbering = doc.part.numbering_part.element
    nums = numbering.findall(qn("w:num"))
    new_id = max((int(n.get(qn("w:numId"))) for n in nums), default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return new_id


def set_paragraph_numbering(paragraph, num_id, ilvl=0):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl_el)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)


# ---------------------------------------------------------------------------
# Inline markdown parsing (`code`, **bold**)
# ---------------------------------------------------------------------------
def _add_run(paragraph, text, bold=False, code=False):
    if not text:
        return
    run = paragraph.add_run(text)
    run.bold = bold
    if code:
        _make_monospace(run)


def add_inline(paragraph, text, bold=False):
    i, n = 0, len(text)
    buffer = ""
    while i < n:
        if text[i] == "`":
            j = text.find("`", i + 1)
            if j == -1:
                buffer += text[i]
                i += 1
                continue
            if buffer:
                _add_run(paragraph, buffer, bold=bold)
                buffer = ""
            _add_run(paragraph, text[i + 1:j], bold=bold, code=True)
            i = j + 1
        elif text.startswith("**", i):
            j = text.find("**", i + 2)
            if j == -1:
                buffer += text[i:i + 2]
                i += 2
                continue
            if buffer:
                _add_run(paragraph, buffer, bold=bold)
                buffer = ""
            add_inline(paragraph, text[i + 2:j], bold=True)
            i = j + 2
        else:
            buffer += text[i]
            i += 1
    if buffer:
        _add_run(paragraph, buffer, bold=bold)


# ---------------------------------------------------------------------------
# Block helpers
# ---------------------------------------------------------------------------
def add_heading(doc, level, text):
    p = doc.add_heading(level=min(level, 9))
    add_inline(p, text)
    return p


def _make_monospace(run):
    run.font.name = CODE_FONT
    run.font.size = CODE_SIZE
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        r_fonts.set(qn(attr), CODE_FONT)


def add_code_block(doc, code_lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Pt(6)
    shade_paragraph(p, CODE_FILL)
    for idx, line in enumerate(code_lines):
        run = p.add_run()
        _make_monospace(run)
        if idx > 0:
            run.add_break()
        run.add_text(line)
    return p


def add_quote(doc, text):
    p = doc.add_paragraph(style="Quote")
    add_inline(p, text)
    return p


def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    add_bottom_border(p)
    return p


def _is_separator_row(line):
    stripped = line.strip()
    if not stripped.startswith("|"):
        return False
    return bool(re.fullmatch(r"\|[\s\-:|]+\|?", stripped)) and "-" in stripped


def _split_row(line):
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.autofit = True
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for c_idx in range(ncols):
            cell = cells[c_idx]
            text = row[c_idx] if c_idx < len(row) else ""
            paragraph = cell.paragraphs[0]
            parts = re.split(r"<br\s*/?>", text, flags=re.I)
            for k, part in enumerate(parts):
                if k > 0:
                    paragraph.add_run().add_break()
                add_inline(paragraph, part, bold=(r_idx == 0))
            if r_idx == 0:
                shade_cell(cell, TABLE_HEADER_FILL)
    doc.add_paragraph()
    return table


# ---------------------------------------------------------------------------
# Conversion
# ---------------------------------------------------------------------------
def build_document(md_text):
    doc = Document()
    lines = md_text.split("\n")
    total = len(lines)
    ordered_num_id = None
    i = 0

    while i < total:
        line = lines[i]
        stripped = line.strip()

        # fenced code block
        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < total and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, code_lines)
            ordered_num_id = None
            continue

        # blank line
        if stripped == "":
            i += 1
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}", stripped) or re.fullmatch(r"\*{3,}", stripped):
            add_hr(doc)
            ordered_num_id = None
            i += 1
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            add_heading(doc, len(m.group(1)), m.group(2).strip())
            ordered_num_id = None
            i += 1
            continue

        # table
        if stripped.startswith("|") and i + 1 < total and _is_separator_row(lines[i + 1]):
            rows = [_split_row(line)]
            i += 2
            while i < total and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            ordered_num_id = None
            continue

        # blockquote
        if stripped.startswith(">"):
            add_quote(doc, stripped[1:].strip())
            ordered_num_id = None
            i += 1
            continue

        # unordered list
        m = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(1).strip())
            ordered_num_id = None
            i += 1
            continue

        # ordered list
        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            if ordered_num_id is None:
                ordered_num_id = create_numbering(doc)
            p = doc.add_paragraph(style="List Number")
            set_paragraph_numbering(p, ordered_num_id)
            add_inline(p, m.group(1).strip())
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        add_inline(p, stripped)
        ordered_num_id = None
        i += 1

    return doc


def convert(md_path, docx_path=None):
    if not os.path.isfile(md_path):
        print(f"[skip] not found: {md_path}")
        return None
    try:
        with open(md_path, "r", encoding="utf-8-sig") as fh:
            text = fh.read()
    except (OSError, UnicodeDecodeError) as exc:
        print(f"[fail] {md_path}: {exc}")
        return None
    if not text.strip():
        print(f"[skip] empty file: {md_path}")
        return None

    if docx_path is None:
        docx_path = os.path.splitext(md_path)[0] + ".docx"
    out_dir = os.path.dirname(os.path.abspath(docx_path))
    os.makedirs(out_dir, exist_ok=True)

    try:
        build_document(text).save(docx_path)
    except (OSError, PermissionError) as exc:
        print(f"[fail] cannot write {docx_path}: {exc}")
        print("       close the file in Word and try again")
        return None
    print(f"[ok]   {md_path}  ->  {docx_path}")
    return docx_path


def collect_inputs(patterns, convert_all, outdir):
    files = []
    if convert_all:
        search_dir = patterns[0] if patterns else "."
        if not os.path.isdir(search_dir):
            search_dir = "."
            files = _expand(patterns)
        for path in sorted(glob.glob(os.path.join(search_dir, "*.md"))):
            if os.path.isfile(path):
                files.append(path)
    else:
        files = _expand(patterns)

    result = []
    for path in files:
        if outdir:
            result.append((path, os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".docx")))
        else:
            result.append((path, None))
    return result


def _expand(patterns):
    found = []
    for pattern in patterns:
        matches = sorted(glob.glob(pattern))
        if matches:
            found.extend(matches)
        elif os.path.isfile(pattern):
            found.append(pattern)
        else:
            print(f"[skip] no such file: {pattern}")
    # keep order, drop duplicates
    seen, unique = set(), []
    for f in found:
        key = os.path.abspath(f)
        if key not in seen:
            seen.add(key)
            unique.append(f)
    return unique


def main(argv=None):
    parser = argparse.ArgumentParser(
        prog="md2docx",
        description="Convert Markdown files to Word (.docx) keeping the text unchanged.",
    )
    parser.add_argument("paths", nargs="*", help=".md files or glob patterns")
    parser.add_argument("-a", "--all", action="store_true",
                        help="convert every .md found in the given folder (default: current folder)")
    parser.add_argument("-o", "--out", default=None,
                        help="output .docx file (only for a single input) or output folder")
    parser.add_argument("-V", "--version", action="version",
                        version=f"%(prog)s {__version__}")
    args = parser.parse_args(argv)

    if not args.paths and not args.all:
        parser.print_help()
        return 1

    jobs = collect_inputs(args.paths, args.all, args.out)
    if not jobs:
        print("Nothing to convert.")
        return 1

    single = len(jobs) == 1
    done = 0
    for src, dst in jobs:
        out = dst
        if args.out and single and args.out.lower().endswith(".docx"):
            out = args.out
        if convert(src, out):
            done += 1
    print(f"Converted: {done}/{len(jobs)}")
    return 0 if done else 1


if __name__ == "__main__":
    sys.exit(main())
