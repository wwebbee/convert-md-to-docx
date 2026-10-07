# -*- coding: utf-8 -*-
"""Tests for md2docx. Stdlib unittest, no extra dependencies.

Run from the repository root:
    python -m unittest discover -s tests -v
"""
import os
import shutil
import sys
import tempfile
import unittest
from io import StringIO
from unittest.mock import patch

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx.oxml.ns import qn  # noqa: E402

import md2docx  # noqa: E402


def texts(document):
    return [p.text for p in document.paragraphs]


def styles(document):
    return [p.style.name for p in document.paragraphs]


def num_id(paragraph):
    if paragraph._p.pPr is None:
        return None
    numpr = paragraph._p.pPr.find(qn("w:numPr"))
    if numpr is None:
        return None
    return numpr.find(qn("w:numId")).get(qn("w:val"))


class InlineTest(unittest.TestCase):
    def test_code_span_is_monospace(self):
        doc = md2docx.build_document("используй `md2docx.py` здесь")
        runs = doc.paragraphs[0].runs
        self.assertEqual("".join(r.text for r in runs), "используй `md2docx.py` здесь".replace("`", ""))
        self.assertTrue(any(r.font.name == md2docx.CODE_FONT for r in runs))

    def test_code_span_content_is_verbatim(self):
        # CommonMark: markup inside backticks is literal, emphasis is not processed.
        doc = md2docx.build_document("# заголовок `**код**`")
        self.assertEqual(doc.paragraphs[0].text, "заголовок **код**")
        code_runs = [r for r in doc.paragraphs[0].runs if r.font.name == md2docx.CODE_FONT]
        self.assertTrue(code_runs)
        self.assertFalse(any(r.bold for r in code_runs))

    def test_unpaired_backtick_is_literal(self):
        doc = md2docx.build_document("обычный ` текст")
        self.assertEqual(doc.paragraphs[0].text, "обычный ` текст")

    def test_unsupported_markup_passes_through(self):
        src = "текст [ссылка](http://example.org) *курсив* ~~зачеркнуто~~"
        doc = md2docx.build_document(src)
        self.assertEqual(doc.paragraphs[0].text, src)


class BlockTest(unittest.TestCase):
    def test_heading_levels(self):
        doc = md2docx.build_document("# a\n\n## b\n\n###### f")
        self.assertEqual(styles(doc)[:3], ["Heading 1", "Heading 2", "Heading 6"])

    def test_code_block_keeps_indentation_and_lines(self):
        doc = md2docx.build_document("```python\nif x:\n    print(1)\n```")
        para = [p for p in doc.paragraphs if p.text][0]
        self.assertEqual(para.text, "if x:\n    print(1)")
        self.assertTrue(all(r.font.name == md2docx.CODE_FONT for r in para.runs))

    def test_fence_content_is_not_parsed(self):
        doc = md2docx.build_document("```\n# не заголовок\n```")
        self.assertEqual([p.text for p in doc.paragraphs if p.text], ["# не заголовок"])

    def test_quote_style(self):
        doc = md2docx.build_document("> цитата")
        self.assertEqual(doc.paragraphs[0].style.name, "Quote")

    def test_horizontal_rule_has_border(self):
        doc = md2docx.build_document("до\n\n---\n\nпосле")
        hr = [p for p in doc.paragraphs if p.text == ""][0]
        self.assertIsNotNone(hr._p.pPr.find(qn("w:pBdr")))

    def test_bullets(self):
        doc = md2docx.build_document("- a\n* b\n+ c")
        self.assertEqual(styles(doc)[:3], ["List Bullet"] * 3)


class ListRestartTest(unittest.TestCase):
    def test_each_ordered_list_gets_its_own_numbering(self):
        doc = md2docx.build_document("1. a\n2. b\n\nабзац\n\n1. c\n2. d")
        numbered = [p for p in doc.paragraphs if p.style.name == "List Number"]
        self.assertEqual(len(numbered), 4)
        self.assertEqual(num_id(numbered[0]), num_id(numbered[1]))
        self.assertNotEqual(num_id(numbered[0]), num_id(numbered[2]))

    def test_restart_override_present(self):
        doc = md2docx.build_document("1. a")
        first = num_id(doc.paragraphs[0])
        numbering = doc.part.numbering_part.element
        nums = [n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == first]
        self.assertEqual(len(nums), 1)
        override = nums[0].find(qn("w:lvlOverride"))
        self.assertEqual(override.find(qn("w:startOverride")).get(qn("w:val")), "1")


class TableTest(unittest.TestCase):
    def test_header_separator_row_is_dropped(self):
        doc = md2docx.build_document("| A | B |\n|:--|--:|\n| 1 | 2 |")
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 2)
        self.assertEqual([c.text for c in table.rows[0].cells], ["A", "B"])

    def test_br_becomes_line_break(self):
        doc = md2docx.build_document("| A |\n|---|\n| раз<br>два |")
        self.assertEqual(doc.tables[0].rows[1].cells[0].text, "раз\nдва")

    def test_ragged_rows_are_padded(self):
        doc = md2docx.build_document("| A | B |\n|---|---|\n| только-одна |")
        self.assertEqual(len(doc.tables[0].columns), 2)
        self.assertEqual(doc.tables[0].rows[1].cells[1].text, "")

    def test_header_cell_is_shaded_and_bold(self):
        doc = md2docx.build_document("| A | B |\n|---|---|\n| 1 | 2 |")
        cell = doc.tables[0].rows[0].cells[0]
        self.assertEqual(cell.paragraphs[0].runs[0].bold, True)
        self.assertIn(qn("w:shd"), [c.tag for c in cell._tc.tcPr])


class NewlineTest(unittest.TestCase):
    def test_crlf_matches_lf(self):
        src = "# Заголовок\n\n- пункт\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"
        lf = md2docx.build_document(src)
        crlf = md2docx.build_document(src.replace("\n", "\r\n"))
        self.assertEqual(texts(lf), texts(crlf))
        self.assertEqual([[c.text for c in r.cells] for r in lf.tables[0].rows],
                         [[c.text for c in r.cells] for r in crlf.tables[0].rows])


class CliTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.src = os.path.join(self.tmp, "in.md")
        with open(self.src, "w", encoding="utf-8") as fh:
            fh.write("#Doc\n")
        # keep the CLI chatter out of the test report
        self._stdout = patch("sys.stdout", new_callable=StringIO)
        self.stdout = self._stdout.start()

    def tearDown(self):
        self._stdout.stop()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_default_output_is_side_by_side(self):
        self.assertEqual(md2docx.convert(self.src), os.path.join(self.tmp, "in.docx"))
        self.assertTrue(os.path.isfile(os.path.join(self.tmp, "in.docx")))

    def test_empty_file_is_skipped(self):
        empty = os.path.join(self.tmp, "empty.md")
        open(empty, "w").close()
        self.assertIsNone(md2docx.convert(empty))

    def test_missing_file_is_skipped(self):
        self.assertIsNone(md2docx.convert(os.path.join(self.tmp, "nope.md")))

    def test_main_with_out_folder(self):
        out = os.path.join(self.tmp, "out")
        self.assertEqual(md2docx.main([self.src, "-o", out]), 0)
        self.assertTrue(os.path.isfile(os.path.join(out, "in.docx")))

    def test_main_with_explicit_docx_name(self):
        out = os.path.join(self.tmp, "itog.docx")
        self.assertEqual(md2docx.main([self.src, "-o", out]), 0)
        self.assertTrue(os.path.isfile(out))

    def test_main_without_arguments_fails(self):
        self.assertEqual(md2docx.main([]), 1)

    def test_main_returns_1_when_nothing_converted(self):
        missing = os.path.join(self.tmp, "missing.md")
        self.assertEqual(md2docx.main([missing]), 1)

    def test_bom_is_stripped(self):
        with open(self.src, "w", encoding="utf-8-sig") as fh:
            fh.write("# Doc\n")
        out = md2docx.convert(self.src)
        from docx import Document
        self.assertEqual(Document(out).paragraphs[0].text, "Doc")


if __name__ == "__main__":
    unittest.main()
